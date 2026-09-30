"""In-memory test documents, generated so no binary fixtures are committed."""

from __future__ import annotations

import io
import zipfile


def make_pdf(pages: list[str], title: str = "Momentum Study") -> bytes:
    """Minimal valid PDF with one line of Helvetica text per page (hand-built, no writer library)."""
    objs: list[bytes] = []
    n_pages = len(pages)
    page_ids = [3 + 2 * i for i in range(n_pages)]
    font_id = 3 + 2 * n_pages
    info_id = font_id + 1
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{p} 0 R" for p in page_ids).encode()
    objs.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(n_pages).encode() + b" >>")
    for i, text in enumerate(pages):
        content_id = page_ids[i] + 1
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode()
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objs.append(f"<< /Title ({title}) /Author (Jane Researcher) /CreationDate (D:20190101000000) >>".encode())

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info {info_id} 0 R >>\nstartxref\n{xref}\n%%EOF\n"
              .encode())
    return out.getvalue()


def make_encrypted_pdf() -> bytes:
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(612, 792)
    w.encrypt("secret-password")
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.core_properties.title = "Overnight Effect Notes"
    d.core_properties.author = "A. Quant"
    d.add_heading("Entry rules", level=1)
    d.add_paragraph("Buy SPY at the close; sell at the next open.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Param", "Value"
    t.cell(1, 0).text, t.cell(1, 1).text = "Lookback", "UNKNOWN"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx(with_chart: bool = True) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "VRP Strategy"
    slide.placeholders[1].text = "Sell 30-delta SPX puts monthly"
    slide.notes_slide.notes_text_frame.text = "Speaker note: see arXiv:1234.56789"
    if with_chart:
        data = CategoryChartData()
        data.categories = ["2019", "2020"]
        data.add_series("Return", (0.1, -0.2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(2), data)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Params"
    ws.append(["name", "value"])
    ws.append(["lookback", 12])
    ws.append(["formula", "=1+1"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_epub(drm: bool = False) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" '
                   'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                   '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                   "</rootfiles></container>")
        z.writestr("OEBPS/content.opf",
                   '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
                   '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Quant Book</dc:title>'
                   "<dc:creator>B. Author</dc:creator><dc:date>2020</dc:date></metadata>"
                   '<manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/></manifest>'
                   '<spine><itemref idref="c1"/></spine></package>')
        z.writestr("OEBPS/ch1.xhtml",
                   "<html><body><h1>Chapter 1: Trend following</h1><p>Use a 200-day moving average.</p>"
                   "</body></html>")
        if drm:
            z.writestr("META-INF/encryption.xml", "<encryption/>")
    return buf.getvalue()


def make_zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return buf.getvalue()
