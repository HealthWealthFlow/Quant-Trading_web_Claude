"""CSV and XLSX. Formulas are never evaluated: XLSX is read with cached values only (spec §7)."""

from __future__ import annotations

import csv
import io

import openpyxl

from ..taxonomy import AccessStatus
from .base import PARSE_ERROR, TRUNCATED, HandlerResult, Location, SourceHandler, Table, TextBlock
from .office import _open_checked
from .text import decode_text

MAX_ROWS = 5000


def _table_text(rows: list[list[str]], limit: int = 200) -> str:
    return "\n".join(" | ".join(r) for r in rows[:limit])


class SpreadsheetHandler(SourceHandler):
    name = "SpreadsheetHandler"
    version = "1"
    formats = ("csv", "xlsx")

    def extract(self, data: bytes, fmt: str, base_url: str | None = None) -> HandlerResult:
        result = self.new_result(data, fmt)
        if fmt == "csv":
            self._csv(data, result)
        else:
            self._xlsx(data, result)
        return result.finalize()

    @staticmethod
    def _csv(data: bytes, result: HandlerResult) -> None:
        text = decode_text(data, result)
        try:
            dialect = csv.Sniffer().sniff(text[:4096])
        except csv.Error:
            dialect = csv.excel
        rows: list[list[str]] = []
        for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
            if i >= MAX_ROWS:
                result.add_limitation(TRUNCATED, f"only first {MAX_ROWS} rows read")
                break
            rows.append([c.strip() for c in row])
        result.tables.append(Table(rows, Location(sheet="csv")))
        result.blocks.append(TextBlock(_table_text(rows), Location(sheet="csv")))

    @staticmethod
    def _xlsx(data: bytes, result: HandlerResult) -> None:
        if not _open_checked(data, result):
            return
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            result.access_status = AccessStatus.ERROR
            result.add_limitation(PARSE_ERROR, str(e)[:200])
            return
        result.set_meta("title", wb.properties.title)
        result.set_meta("author", wb.properties.creator)
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    result.add_limitation(TRUNCATED, f"sheet {ws.title}: only first {MAX_ROWS} rows read")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    rows.append(cells)
            if rows:
                loc = Location(sheet=ws.title)
                result.tables.append(Table(rows, loc))
                result.blocks.append(TextBlock(_table_text(rows), loc))
        wb.close()
