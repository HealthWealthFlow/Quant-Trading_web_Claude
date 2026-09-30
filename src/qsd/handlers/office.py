"""Word (DOCX) and PowerPoint (PPTX). Containers are safety-checked first; macros are never run (spec §7, §22)."""

from __future__ import annotations

import io

import docx
from pptx import Presentation

from ..taxonomy import AccessStatus
from .archive import UnsafeArchiveError, assert_safe_zip
from .base import (
    MACRO_CONTAINER,
    PARSE_ERROR,
    UNREADABLE_CHART_DATA,
    UNSAFE_ARCHIVE,
    HandlerResult,
    Location,
    SourceHandler,
    Table,
    TextBlock,
)


def _open_checked(data: bytes, result: HandlerResult) -> bool:
    try:
        z = assert_safe_zip(data)
    except UnsafeArchiveError as e:
        result.access_status = AccessStatus.ERROR
        result.add_limitation(UNSAFE_ARCHIVE, str(e))
        return False
    if any(n.lower().endswith("vbaproject.bin") for n in z.names()):
        result.add_limitation(MACRO_CONTAINER, "document contains macros; parsed as data only")
    return True


def _core_meta(props: object, result: HandlerResult) -> None:
    result.set_meta("title", getattr(props, "title", None))
    result.set_meta("author", getattr(props, "author", None))
    created = getattr(props, "created", None)
    if created:
        # A file property, often from a template: never treated as the publication date (spec §1).
        result.set_meta("file_created_date", created.date().isoformat())


class DOCXHandler(SourceHandler):
    name = "DOCXHandler"
    version = "1"
    formats = ("docx",)

    def extract(self, data: bytes, fmt: str, base_url: str | None = None) -> HandlerResult:
        result = self.new_result(data, fmt)
        if not _open_checked(data, result):
            return result.finalize()
        try:
            d = docx.Document(io.BytesIO(data))
        except Exception as e:  # noqa: BLE001
            result.access_status = AccessStatus.ERROR
            result.add_limitation(PARSE_ERROR, str(e)[:200])
            return result.finalize()
        _core_meta(d.core_properties, result)
        section: str | None = None
        for p in d.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith(("heading", "title")):
                section = text[:200]
            result.blocks.append(TextBlock(text, Location(section=section)))
        for i, t in enumerate(d.tables, 1):
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            result.tables.append(Table(rows, Location(item=f"table {i}")))
        for rel in d.part.rels.values():
            if "hyperlink" in rel.reltype and rel.is_external:
                result.links.append(rel.target_ref)
        return result.finalize()


class PPTXHandler(SourceHandler):
    name = "PPTXHandler"
    version = "1"
    formats = ("pptx",)

    def extract(self, data: bytes, fmt: str, base_url: str | None = None) -> HandlerResult:
        result = self.new_result(data, fmt)
        if not _open_checked(data, result):
            return result.finalize()
        try:
            prs = Presentation(io.BytesIO(data))
        except Exception as e:  # noqa: BLE001
            result.access_status = AccessStatus.ERROR
            result.add_limitation(PARSE_ERROR, str(e)[:200])
            return result.finalize()
        _core_meta(prs.core_properties, result)
        for n, slide in enumerate(prs.slides, 1):
            title_shape = slide.shapes.title
            title = title_shape.text_frame.text.strip() if title_shape is not None else None
            for shape in slide.shapes:
                self._shape(shape, n, title, result)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    result.blocks.append(TextBlock(notes, Location(slide=n, section=title, item="notes")))
        return result.finalize()

    @staticmethod
    def _shape(shape: object, n: int, title: str | None, result: HandlerResult) -> None:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                result.blocks.append(TextBlock(text, Location(slide=n, section=title)))
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.hyperlink and run.hyperlink.address:
                        result.links.append(run.hyperlink.address)
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
            result.tables.append(Table(rows, Location(slide=n, section=title, item="table")))
        if getattr(shape, "has_chart", False) and shape.has_chart:
            # Chart values are not read: cached chart data can be stale or partial (spec §22: never invent).
            result.add_limitation(UNREADABLE_CHART_DATA, f"slide {n}")
        if hasattr(shape, "shapes"):  # group shape
            for child in shape.shapes:
                PPTXHandler._shape(child, n, title, result)
